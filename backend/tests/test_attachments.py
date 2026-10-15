import asyncio
import io
import zipfile
from uuid import uuid4

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app import attachments
from app.main import app
from test_api import auth, register


def archive_bytes(parts):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts:
            archive.writestr(name, data)
    return buffer.getvalue()


@pytest.mark.parametrize("budget", ["xml", "total", "count"])
def test_docx_expansion_rejected_before_document_parser(monkeypatch, budget):
    import docx
    def must_not_parse(*args):
        pytest.fail("oversized archive reached the document parser")
    monkeypatch.setattr(docx, "Document", must_not_parse)
    if budget == "xml":
        parts = [("word/document.xml", b" " * (attachments.MAX_DOCX_XML_BYTES + 1))]
    elif budget == "total":
        monkeypatch.setattr(attachments, "MAX_DOCX_EXPANDED_BYTES", 1024)
        parts = [(f"part{i}.bin", b"x" * 512) for i in range(3)]
    else:
        monkeypatch.setattr(attachments, "MAX_DOCX_PARTS", 3)
        parts = [(f"part{i}.xml", b"<x/>") for i in range(4)]
    packed = archive_bytes(parts)
    assert len(packed) < 100_000
    with pytest.raises(HTTPException) as error:
        attachments.extract_file("synthetic.docx", packed)
    assert error.value.status_code == 413


@pytest.mark.parametrize("extension", ["txt", "docx", "pdf"])
def test_supported_uploads_use_bounded_parser_and_preserve_text(extension):
    data = "本人对甘草过敏\n普通病历".encode()
    if extension == "docx":
        document = Document()
        document.add_paragraph("普通病历")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "本人对甘草过敏"
        buffer = io.BytesIO()
        document.save(buffer)
        data = buffer.getvalue()
    elif extension == "pdf":
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        buffer = io.BytesIO()
        writer.write(buffer)
        data = buffer.getvalue()
    async def exercise():
        slots = asyncio.Semaphore(2)
        kind, content = await attachments.parse_attachment("synthetic." + extension, data, slots)
        assert kind
        assert not slots.locked()
        if extension != "pdf":
            assert "本人对甘草过敏" in content
            assert "普通病历" in content
    asyncio.run(exercise())


def test_pdf_page_limit_is_enforced_in_worker():
    writer = PdfWriter()
    for _ in range(attachments.MAX_PDF_PAGES + 1):
        writer.add_blank_page(width=100, height=100)
    buffer = io.BytesIO()
    writer.write(buffer)
    async def exercise():
        with pytest.raises(HTTPException) as error:
            await attachments.parse_attachment("large.pdf", buffer.getvalue(), asyncio.Semaphore(1))
        assert error.value.status_code == 413
    asyncio.run(exercise())


@pytest.mark.parametrize("stop", ["timeout", "cancel"])
def test_worker_is_killed_reaped_and_slot_released(monkeypatch, stop):
    async def exercise():
        entered = asyncio.Event()
        class Worker:
            returncode = None
            killed = False
            reaped = False
            async def communicate(self, data):
                entered.set()
                await asyncio.Future()
            def kill(self):
                self.killed = True
                self.returncode = -9
            async def wait(self):
                self.reaped = True
        worker = Worker()
        async def spawn(*args, **kwargs):
            assert kwargs["stderr"] == asyncio.subprocess.DEVNULL
            return worker
        monkeypatch.setattr(asyncio, "create_subprocess_exec", spawn)
        monkeypatch.setattr(attachments, "PARSER_TIMEOUT_SECONDS", .02)
        slots = asyncio.Semaphore(1)
        task = asyncio.create_task(attachments.parse_attachment("slow.pdf", b"%PDF", slots))
        await entered.wait()
        if stop == "cancel":
            task.cancel()
            with pytest.raises(asyncio.CancelledError):
                await task
        else:
            with pytest.raises(HTTPException) as error:
                await task
            assert error.value.status_code == 413
        assert worker.killed and worker.reaped
        assert not slots.locked()
    asyncio.run(exercise())


def test_busy_parser_rejects_without_launching_more_workers(monkeypatch):
    async def must_not_spawn(*args, **kwargs):
        pytest.fail("parser concurrency limit was bypassed")
    monkeypatch.setattr(asyncio, "create_subprocess_exec", must_not_spawn)
    async def exercise():
        with pytest.raises(HTTPException) as error:
            await attachments.parse_attachment("note.pdf", b"%PDF", asyncio.Semaphore(0))
        assert error.value.status_code == 429
    asyncio.run(exercise())


def test_rejected_docx_does_not_create_upload_or_store_file():
    from app.config import UPLOAD_DIR
    packed = archive_bytes([("word/document.xml", b" " * (attachments.MAX_DOCX_XML_BYTES + 1))])
    with TestClient(app) as client:
        token, _ = register(client, "archive_" + uuid4().hex[:8])
        before = set(UPLOAD_DIR.iterdir())
        response = client.post("/api/upload", headers=auth(token), files={"file": ("large.docx", packed)})
        assert response.status_code == 413
        assert client.get("/api/uploads", headers=auth(token)).json() == []
        assert set(UPLOAD_DIR.iterdir()) == before
