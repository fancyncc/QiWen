import asyncio
import io
import json
import threading
from uuid import uuid4

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient

from app import knowledge, model, workflow
from app.agents import tools
from app.agents.runtime import EmptyArgs, ToolSpec, _execute_tool
from app.documents import MAX_FILE_TEXT_CHARS, model_excerpt
from app.main import app, extract_file
from test_api import auth, register


def test_sync_tool_timeout_does_not_block_loop_or_return_late_result():
    async def scenario():
        release, entered = threading.Event(), threading.Event()
        def blocking(*_):
            entered.set()
            release.wait(2)
            return "late result"
        task = asyncio.create_task(_execute_tool(ToolSpec("slow", "synthetic", EmptyArgs, blocking, .03), {}, {}))
        try:
            await asyncio.sleep(.06)
            assert entered.is_set()
            assert task.done(), "event loop must deliver the timeout before the worker finishes"
            with pytest.raises(TimeoutError):
                await task
        finally:
            release.set()
    asyncio.run(scenario())


def test_attachment_tail_and_newline_boundaries_are_preserved():
    text = "普通记录。\n" * 900 + "我最近失眠三天。\n我现在呼吸困难。"
    parsed = workflow.document_parse({"file_text": text})["file_text"]
    assert parsed == text
    assert "呼吸困难" in model_excerpt(parsed)


@pytest.mark.parametrize("extension", ["txt", "docx"])
def test_oversized_attachment_is_rejected_instead_of_truncated(extension):
    text = "文" * (MAX_FILE_TEXT_CHARS + 1)
    data = text.encode()
    if extension == "docx":
        doc = Document()
        doc.add_paragraph(text)
        buffer = io.BytesIO()
        doc.save(buffer)
        data = buffer.getvalue()
    with pytest.raises(HTTPException) as error:
        extract_file("large." + extension, data)
    assert error.value.status_code == 413


def test_docx_tables_are_included_in_extraction():
    doc = Document()
    doc.add_paragraph("病历")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "对甘草过敏"
    buffer = io.BytesIO()
    doc.save(buffer)
    assert "对甘草过敏" in extract_file("table.docx", buffer.getvalue())[1]


def test_second_long_attachment_tail_is_checked_by_workflow():
    with TestClient(app) as client:
        token, _ = register(client, "attach_" + uuid4().hex[:8])
        headers = auth(token)
        files = []
        for index, text in enumerate(["普通记录。\n" * 3000, "普通记录。\n" * 2000 + "现在呼吸困难。"]):
            uploaded = client.post("/api/upload", headers=headers,
                                   files={"file": (f"file{index}.txt", text.encode(), "text/plain")})
            assert uploaded.status_code == 201
            files.append(uploaded.json()["id"])
        result = client.post("/api/chat", headers=headers, json={"message": "请看附件", "files": files})
        assert result.status_code == 200
        data = result.json()
        assert data["intent_type"] == "safety_rejected"
        assert "急救" in data["response"]
        assert not any(event["stage"] == "formula" for event in data["trace"]["events"])


def test_late_attachment_symptoms_survive_ownership_confirmation():
    with TestClient(app) as client:
        token, _ = register(client, "tail_" + uuid4().hex[:8])
        headers = auth(token)
        text = "普通记录。\n" * 1000 + "失眠三天，白天疲劳。"
        upload = client.post("/api/upload", headers=headers, files={"file": ("late.txt", text.encode(), "text/plain")}).json()
        first = client.post("/api/chat", headers=headers, json={"message": "请看附件", "files": [upload["id"]]}).json()
        second = client.post("/api/chat", headers=headers,
                             json={"message": "附件是我本人的病历", "thread_id": first["thread_id"]}).json()
        assert {"失眠", "疲劳"} <= set(second["symptoms"])


def test_knowledge_model_receives_only_public_evidence(monkeypatch):
    monkeypatch.setattr(model, "MODEL_MODE", "openai_compatible")
    async def answer(system, user, *, role, schema):
        data = json.loads(user)
        assert role == "knowledge"
        assert set(data) == {"question", "evidence"}
        assert "private-marker" not in user
        assert data["evidence"][0]["id"] == "S1"
        return {"answer": "阴阳用于描述相互联系的两个方面。", "source_ids": ["S1"]}
    monkeypatch.setattr(model, "acomplete_json", answer)
    state = {"message": "什么是阴阳？", "profile": {"history": "private-marker"}, "memory_context": "private-marker",
             "docs": [{"id": "source-one", "title": "合成资料", "chapter": "概念", "body": "阴阳是相互联系的两个方面。"}]}
    result = asyncio.run(knowledge.answer_knowledge(state))
    assert "相互联系" in result["response"] and "合成资料" in result["response"]


@pytest.mark.parametrize("invalid", ["source", "dosage"])
def test_knowledge_rejects_unknown_sources_and_dosing(monkeypatch, invalid):
    monkeypatch.setattr(model, "MODEL_MODE", "openai_compatible")
    async def answer(*args, **kwargs):
        return {"answer": "甘草 10g" if invalid == "dosage" else "概念解释", "source_ids": ["invented"] if invalid == "source" else []}
    monkeypatch.setattr(model, "acomplete_json", answer)
    if invalid == "source":
        with pytest.raises(RuntimeError):
            asyncio.run(knowledge.answer_knowledge({"message": "什么是气血？"}))
    else:
        assert "10g" not in asyncio.run(knowledge.answer_knowledge({"message": "什么是气血？"}))["response"]


@pytest.mark.parametrize("mode", ["fixed", "multi"])
def test_knowledge_graph_uses_public_branch_and_preserves_clinical_state(monkeypatch, mode):
    async def retrieve(plan):
        assert plan["lexical_terms"]
        return [], {"warning": "", "selected": []}
    monkeypatch.setattr(knowledge, "retrieve_classics", retrieve)
    with TestClient(app) as client:
        token, _ = register(client, "know_" + uuid4().hex[:8])
        headers = auth(token)
        thread = client.post("/api/threads", headers=headers, json={"mode": mode}).json()["id"]
        data = client.post("/api/chat", headers=headers, json={"thread_id": thread, "message": "什么是阴阳？"}).json()
        stages = {event["stage"] for event in data["trace"]["events"]}
        assert {"knowledge_search", "knowledge_answer"} <= stages
        assert not stages & {"formula", "bianzheng", "herb"}
        assert data["trace"]["identity_mode"] == "knowledge"
        assert "什么是阴阳" in data["response"] and "Mock" in data["response"]
