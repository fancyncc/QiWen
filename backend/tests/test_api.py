from __future__ import annotations

import io
import json
import os
import tempfile
from uuid import uuid4

os.environ["QIWEN_DATA_DIR"] = tempfile.mkdtemp(prefix="qiwen-test-")

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app.main import app
from app.retrieval import ingest_file, retrieve
from app.safety import check_herbs


def register(client: TestClient, name: str) -> tuple[str, str]:
    result = client.post("/api/register", json={"username": name, "password": "correct-horse-123"})
    assert result.status_code == 201, result.text
    data = result.json()
    return data["access_token"], data["user"]["id"]


def auth(token: str) -> dict:
    return {"Authorization": f"Bearer {token}"}


def test_accounts_threads_profiles_and_file_isolation():
    with TestClient(app) as client:
        a, a_id = register(client, "usera_" + uuid4().hex[:6])
        b, b_id = register(client, "userb_" + uuid4().hex[:6])
        thread = client.post("/api/threads", headers=auth(a), json={"mode": "fixed"}).json()["id"]
        assert client.get(f"/api/threads/{thread}/messages", headers=auth(b)).status_code == 404
        assert client.put(f"/api/threads/{thread}", headers=auth(b), json={"title": "steal"}).status_code == 404
        assert client.delete(f"/api/threads/{thread}", headers=auth(b)).status_code == 404
        assert client.get(f"/api/profiles/{a_id}", headers=auth(b)).status_code == 403
        assert client.patch(f"/api/profiles/{a_id}", headers=auth(b), json={"allergies": "甘草"}).status_code == 403
        assert client.patch(f"/api/profiles/{a_id}", headers=auth(a), json={"allergies": "甘草"}).status_code == 200
        assert client.get(f"/api/profiles/{a_id}", headers=auth(a)).json()["allergies"] == "甘草"
        uploaded = client.post("/api/upload", headers=auth(a), files={"file": ("note.txt", b"hello", "text/plain")})
        assert uploaded.status_code == 201
        file_id = uploaded.json()["id"]
        assert any(item["id"] == file_id for item in client.get("/api/uploads", headers=auth(a)).json())
        assert all(item["id"] != file_id for item in client.get("/api/uploads", headers=auth(b)).json())
        denied = client.post("/api/chat", headers=auth(b), json={"message": "你好", "files": [file_id]})
        assert denied.status_code == 404
        assert client.delete(f"/api/uploads/{file_id}", headers=auth(b)).status_code == 404
        assert client.delete(f"/api/uploads/{file_id}", headers=auth(a)).status_code == 204
        assert all(item["id"] != file_id for item in client.get("/api/uploads", headers=auth(a)).json())
        assert client.post("/api/chat", headers=auth(b), json={"message": "你好", "thread_id": thread}).status_code == 404
        assert client.get("/api/profiles/" + b_id, headers=auth(b)).status_code == 200


def test_followup_sse_multi_and_history():
    with TestClient(app) as client:
        token, user_id = register(client, "flow_" + uuid4().hex[:6])
        headers = auth(token)
        thread = client.post("/api/threads", headers=headers, json={"mode": "fixed"}).json()["id"]
        first = client.post("/api/chat", headers=headers, json={"message": "我最近失眠", "thread_id": thread}).json()
        assert first.get("intent_type") == "consult", first
        assert client.get("/api/threads", headers=headers).json()[0]["title"] == "我最近失眠"
        second = client.post("/api/chat", headers=headers, json={"message": "三天，白天有些疲劳", "thread_id": thread}).json()
        assert second["intent_type"] == "diagnose"
        assert "失眠" in second["response"]
        assert len(client.get(f"/api/threads/{thread}/messages", headers=headers).json()) == 4
        assert "甘草" in second["response"]
        client.patch(f"/api/profiles/{user_id}", headers=headers, json={"allergies": "甘草"})
        blocked = client.post("/api/chat", headers=headers, json={"message": "我最近失眠三天，伴有疲劳", "thread_id": thread}).json()
        assert "药材清单已隐藏" in blocked["response"]
        assert "酸枣仁、茯苓" not in blocked["response"]

        multi = client.post("/api/threads", headers=headers, json={"mode": "multi"}).json()["id"]
        with client.stream("POST", "/api/chat?stream=true", headers=headers,
                           json={"message": "我最近咳嗽三天，伴有疲劳", "thread_id": multi}) as response:
            events = [json.loads(line[6:]) for line in response.iter_lines() if line.startswith("data: ")]
        assert response.status_code == 200
        assert any(item.get("stage") == "herb" for item in events)
        assert {item["payload"].get("agent") for item in events if item["type"] == "agent.action"} == {
            "辨证 Agent", "方剂 Agent", "本草 Agent"
        }
        def at(stage: str, status: str) -> int:
            return next(i for i, event in enumerate(events)
                        if event.get("stage") == stage and event.get("status") == status)
        for stage in ("inquiry", "bianzheng", "formula", "herb", "synthesis", "reviewer", "safety"):
            assert at(stage, "start") < at(stage, "done")
        assert at("synthesis", "start") > max(at(stage, "done") for stage in ("bianzheng", "formula", "herb"))
        assert any(item["type"] == "candidate.delta" for item in events)
        assert events[-1]["done"] is True
        assert set(events[-1]["trace"]["agents"]["artifacts"]) == {"bianzheng", "formula", "herb"}
        assert len(events[-1]["notes"]) == 3
        assert [note["expert"] for note in events[-1]["notes"]] == ["辨证", "方剂", "本草"]
        assert all(note.get("evidence") and note.get("reasoning") for note in events[-1]["notes"])
        assert "药材清单已隐藏" in events[-1]["notes"][1]["content"]
        assert events[-1]["notes"][2]["warning"]
        assert len(client.get(f"/api/threads/{multi}/messages", headers=headers).json()) == 2

        client.patch(f"/api/profiles/{user_id}", headers=headers, json={"allergies": ""})
        example = client.post("/api/chat", headers=headers,
                              json={"message": "我最近总是睡不好，白天容易疲劳"}).json()
        assert "失眠" in example["symptoms"]
        assert "酸枣仁汤" in example["response"]


def test_multi_expert_cards_are_persisted_after_safety_review():
    with TestClient(app) as client:
        token, user_id = register(client, "experts_" + uuid4().hex[:6])
        headers = auth(token)
        thread = client.post("/api/threads", headers=headers, json={"mode": "multi"}).json()["id"]
        result = client.post("/api/chat", headers=headers,
                             json={"message": "我最近失眠三天，白天疲劳", "thread_id": thread}).json()
        assert "## 会诊综合" in result["response"]
        assert [note["expert"] for note in result["notes"]] == ["辨证", "方剂", "本草"]
        assert result["notes"][0]["status"] == "资料有限"
        assert "酸枣仁汤" in result["notes"][1]["content"]
        assert result["notes"][2]["warning"] == ""
        saved = client.get(f"/api/threads/{thread}/messages", headers=headers).json()[-1]
        assert saved["notes"] == result["notes"]

        client.patch(f"/api/profiles/{user_id}", headers=headers, json={"allergies": "甘草"})
        blocked = client.post("/api/chat", headers=headers,
                              json={"message": "我最近失眠三天，白天疲劳", "thread_id": thread}).json()
        assert "药材清单已隐藏" in blocked["response"]
        assert "酸枣仁汤" not in blocked["notes"][1]["content"]
        assert "甘草" in blocked["notes"][2]["warning"]


def test_contextual_followup_preserves_denials_without_searching_them(monkeypatch):
    from app import workflow
    from app.agents import tools as agent_tools

    assert workflow.symptom_observations("没有口干口苦、腹痛，舌苔发白") == (
        ["舌苔白"], ["腹痛", "口干", "口苦"],
    )
    queries = []
    async def capture(plan):
        queries.append(plan["keyword_query"])
        return [], {"warning": ""}
    monkeypatch.setattr(agent_tools, "retrieve_classics", capture)
    with TestClient(app) as client:
        token, _ = register(client, "followup_" + uuid4().hex[:6])
        headers = auth(token)
        thread = client.post("/api/threads", headers=headers, json={"mode": "multi"}).json()["id"]
        first = client.post("/api/chat", headers=headers,
                            json={"message": "我肚子胀、怕冷、大便稀", "thread_id": thread}).json()
        assert first["intent_type"] == "consult"
        assert "腹痛" in first["response"] and "舌苔" in first["response"]
        assert "formula" not in [event["stage"] for event in
               client.get(f"/api/threads/{thread}/messages", headers=headers).json()[-1]["trace"]["events"]]

        second = client.post("/api/chat", headers=headers,
                             json={"message": "没有口干口苦、腹痛，舌苔发白，大概两三天吧", "thread_id": thread}).json()
        assert second["intent_type"] == "diagnose"
        assert {"腹胀", "怕冷", "大便稀", "舌苔白"} <= set(second["symptoms"])
        assert not {"腹痛", "口干", "口苦"} & set(second["symptoms"])
        assert "用户表示没有" in second["response"]
        assert "腹痛" in second["notes"][0]["evidence"]
        assert [note["expert"] for note in second["notes"]] == ["辨证", "方剂", "本草"]
        assert queries and all("腹痛" not in query and "口干" not in query and "口苦" not in query
                               for query in queries)


def test_upload_validation_and_retrieval(tmp_path, monkeypatch):
    with TestClient(app) as client:
        token, _ = register(client, "files_" + uuid4().hex[:6])
        headers = auth(token)
        assert client.post("/api/upload", headers=headers, files={"file": ("bad.exe", b"bad")}).status_code == 400
        assert client.post("/api/upload", headers=headers, files={"file": ("huge.txt", b"x" * (10 * 1024 * 1024 + 1))}).status_code == 413
        doc = Document()
        doc.add_paragraph("测试病历内容")
        stream = io.BytesIO()
        doc.save(stream)
        assert client.post("/api/upload", headers=headers, files={"file": ("note.docx", stream.getvalue())}).status_code == 201
        symptom_doc = Document()
        symptom_doc.add_paragraph("已经失眠三天，伴有疲劳。")
        symptom_stream = io.BytesIO()
        symptom_doc.save(symptom_stream)
        symptom_file = client.post("/api/upload", headers=headers,
                                  files={"file": ("symptoms.docx", symptom_stream.getvalue())}).json()["id"]
        attached = client.post("/api/chat", headers=headers,
                               json={"message": "请看附件", "files": [symptom_file]}).json()
        assert attached["intent_type"] == "safety_advisory"
        assert attached["trace"]["safety"]["pending"]["key"] == "subject"
        attached_trace = client.get(f'/api/threads/{attached["thread_id"]}/messages', headers=headers).json()[-1]["trace"]["events"]
        assert any(event["stage"] == "document_parse" and event["status"] == "done" for event in attached_trace)
        confirmed = client.post("/api/chat", headers=headers,
                                json={"message": "附件是我本人的病历", "thread_id": attached["thread_id"]}).json()
        assert confirmed["intent_type"] == "diagnose"
        assert "失眠" in confirmed["symptoms"]
        assert all(f["source"]["kind"] == "attachment" and f.get("ownership_confirmed_by")
                   for f in confirmed["trace"]["inquiry"]["facts"])
        writer = PdfWriter()
        writer.add_blank_page(width=300, height=300)
        valid_pdf = io.BytesIO()
        writer.write(valid_pdf)
        assert client.post("/api/upload", headers=headers, files={"file": ("blank.pdf", valid_pdf.getvalue())}).status_code == 201
        pdf = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
        assert client.post("/api/upload", headers=headers, files={"file": ("broken.pdf", pdf)}).status_code == 400

    text_file = tmp_path / "book.txt"
    text_file.write_text("桂枝汤相关的古籍学习片段。这里是用于检索测试的文本内容。", encoding="utf-8")
    assert ingest_file(text_file, "测试古籍") == 1
    assert ingest_file(text_file, "测试古籍") == 0
    assert retrieve("用于检索测试的文本内容")[0]["title"] == "测试古籍"
    from app import retrieval
    monkeypatch.setattr(retrieval, "RAG_MODE", "hybrid")
    def unavailable(_texts):
        raise RuntimeError("offline")
    monkeypatch.setattr(retrieval, "_embed", unavailable)
    fallback, warning = retrieval.retrieve_with_status("桂枝汤")
    assert fallback and "退化为关键词检索" in warning


def test_contraindication_rules():
    assert check_herbs(["炙甘草", "甘遂"])
    assert check_herbs(["生川乌", "法半夏"])
    assert check_herbs(["官桂", "赤石脂"])
    assert check_herbs(["硫黄", "朴硝"])
    assert check_herbs(["未知药材"])
    assert check_herbs([])
    assert not check_herbs(["酸枣仁", "茯苓", "知母", "川芎", "甘草"])


def test_benchmark_safety_long_tails_and_short_complaints():
    blocked = ["川乌能配半夏吗？", "夹竹桃叶煮水当偏方喝可以吗？",
               "刚才大量咯血，现在很虚弱怎么办？", "腹部硬得像板并剧痛，能否先喝中药？"]
    consult = ["我这两天咳嗽", "我最近胃不舒服", "我腰有点酸", "我晚上容易出汗"]
    with TestClient(app) as client:
        for index, message in enumerate(blocked):
            token, _ = register(client, f"longtail_{index}_{uuid4().hex[:5]}")
            data = client.post("/api/chat", headers=auth(token), json={"message": message}).json()
            assert data["intent_type"] in {"safety_rejected", "safety_advisory"}
            assert "方剂知识示例" not in data["response"]
        for index, message in enumerate(consult):
            token, _ = register(client, f"short_{index}_{uuid4().hex[:5]}")
            data = client.post("/api/chat", headers=auth(token), json={"message": message}).json()
            assert data["intent_type"] == "consult"
        token, _ = register(client, "detailed_" + uuid4().hex[:5])
        detailed = client.post("/api/chat", headers=auth(token),
                               json={"message": "我胸闷一个月，常叹气，压力大时加重。"}).json()
        assert detailed["intent_type"] == "diagnose"


def test_followup_herb_addition_checks_prior_formula_and_isolates_threads():
    from app.db import connection

    with TestClient(app) as client:
        token, _ = register(client, "herb_" + uuid4().hex[:6])
        headers = auth(token)
        original = client.post("/api/chat", headers=headers,
                               json={"message": "我最近失眠三天，白天疲劳"}).json()
        thread = original["thread_id"]
        assert "甘草" in original["response"]
        warning = client.post("/api/chat", headers=headers,
                              json={"message": "可以加海藻吗？", "thread_id": thread}).json()
        assert warning["intent_type"] == "safety_rejected"
        assert all(word in warning["response"] for word in ("安全警告", "甘草", "海藻", "十八反"))
        assert "严重毒副作用" not in warning["response"]
        trace = client.get(f"/api/threads/{thread}/messages", headers=headers).json()[-1]["trace"]["events"]
        stages = [event["stage"] for event in trace if event["status"] == "start"]
        assert "pre_safety" in stages and "inquiry" not in stages and "formula" not in stages

        independent = client.post("/api/chat", headers=headers,
                                  json={"message": "可以加海藻吗？"}).json()
        assert independent["intent_type"] == "safety_advisory"
        assert "完整药材清单" in independent["response"]
        assert "甘草" not in independent["response"]

        with connection() as db:
            db.execute("UPDATE messages SET plan=NULL WHERE thread_id=? AND role='assistant'", (thread,))
        legacy = client.post("/api/chat", headers=headers,
                             json={"message": "海藻能和这个方子一起用吗？", "thread_id": thread}).json()
        assert legacy["intent_type"] == "safety_advisory"
        assert "完整药材清单" in legacy["response"]  # A legacy visible example has no patient-bound structured candidate.

        direct = client.post("/api/chat", headers=headers,
                             json={"message": "甘草和海藻可以一起用吗？"}).json()
        assert direct["intent_type"] == "safety_rejected"
        assert "十八反·甘草" in direct["response"]


def test_bundled_classics_seed_is_complete_and_idempotent():
    from app.corpus_seed import BOOKS, seed_bundled_corpus
    from app.db import connection

    with TestClient(app):
        pass
    with connection() as db:
        counts = dict(db.execute("SELECT title, COUNT(*) FROM corpus GROUP BY title"))
        assert all(counts.get(book, 0) > 0 for book in BOOKS)
        assert sum(counts[book] for book in BOOKS) >= 1500
        assert seed_bundled_corpus(db) == 0


def test_bundled_data_can_be_rebuilt_from_preserved_sources():
    from data_pipeline.ingest import chunk_all_books
    from scripts.rebuild_classics import CHUNK_DIR, KNOWLEDGE_FILE, RAW_DIR, rebuild_knowledge

    generated = chunk_all_books(str(RAW_DIR))
    assert sum(map(len, generated.values())) == 1515
    for book, chunks in generated.items():
        saved = json.loads((CHUNK_DIR / f"{book}.json").read_text(encoding="utf-8"))
        assert saved == [
            {"id": chunk.id, "content": chunk.content, "metadata": chunk.metadata}
            for chunk in chunks
        ]
    saved_knowledge = json.loads(KNOWLEDGE_FILE.read_text(encoding="utf-8"))
    assert saved_knowledge == rebuild_knowledge()
    assert [len(saved_knowledge[key]) for key in ("herbs", "formulas", "contraindications")] == [33, 22, 14]


def test_full_branching_flow_persists_trace_and_session_summary(monkeypatch):
    from app import workflow

    calls = []
    async def staged_retrieval(plan):
        calls.append(plan["keyword_query"])
        assert plan["expansions"] and "不寐" in plan["keyword_query"]
        return [{"title": "测试古籍", "chapter": "测试章", "body": "失眠疲劳的测试片段。"}], {"warning": ""}
    monkeypatch.setattr(workflow, "retrieve_classics", staged_retrieval)
    with TestClient(app) as client:
        token, user_id = register(client, "branch_" + uuid4().hex[:6])
        other, _ = register(client, "branch_" + uuid4().hex[:6])
        headers = auth(token)
        client.patch(f"/api/profiles/{user_id}", headers=headers, json={"allergies": "甘草"})
        thread_id = client.post("/api/threads", headers=headers, json={"mode": "fixed"}).json()["id"]
        result = client.post("/api/chat", headers=headers,
                             json={"message": "我最近失眠三天，白天疲劳", "thread_id": thread_id})
        assert result.status_code == 200, result.text
        assert "药材清单已隐藏" in result.json()["response"]
        assert len(calls) == 1
        saved = client.get(f"/api/threads/{thread_id}/messages", headers=headers).json()[-1]
        stages = [item["stage"] for item in saved["trace"]["events"] if item["status"] == "start"]
        assert stages[:7] == ["prepare", "document_parse", "subject_resolve", "episode_resolve", "profile_load", "pre_safety", "inquiry"]
        assert stages.count("search") == 1
        assert stages.count("revise_formula") == 2
        assert stages.count("safety") == 3
        assert "rewrite_query" not in stages and "evaluate_retrieval" in stages
        assert stages[-3:] == ["context_budget", "episode_write", "finish"]
        profile = client.get(f"/api/profiles/{user_id}", headers=headers).json()
        assert profile["allergies"] == "甘草"
        from app.db import one
        assert "memories" not in profile
        assert one("SELECT summary FROM consultation_episodes WHERE thread_id=?", (thread_id,))
        assert not one("SELECT * FROM profile_memories WHERE thread_id=?", (thread_id,))
        assert client.delete(f"/api/profiles/{user_id}/memories/{thread_id}", headers=auth(other)).status_code == 403
        assert client.delete(f"/api/profiles/{user_id}/memories/{thread_id}", headers=headers).status_code == 204
        assert "memories" not in client.get(f"/api/profiles/{user_id}", headers=headers).json()


def test_skip_urgent_and_exhausted_retrieval_routes(monkeypatch):
    from app import workflow

    with TestClient(app) as client:
        token, _ = register(client, "routes_" + uuid4().hex[:6])
        headers = auth(token)
        skipped = client.post("/api/chat", headers=headers,
                              json={"message": "我最近失眠三天，不用查古籍"}).json()
        skipped_trace = client.get(f'/api/threads/{skipped["thread_id"]}/messages', headers=headers).json()[-1]["trace"]["events"]
        skipped_stages = [event["stage"] for event in skipped_trace if event["status"] == "start"]
        assert "rag_decision" in skipped_stages and "search" not in skipped_stages
        assert "本轮按要求未检索古籍" in skipped["response"]

        urgent = client.post("/api/chat", headers=headers,
                             json={"message": "现在胸痛，呼吸困难"}).json()
        urgent_trace = client.get(f'/api/threads/{urgent["thread_id"]}/messages', headers=headers).json()[-1]["trace"]["events"]
        urgent_stages = [event["stage"] for event in urgent_trace if event["status"] == "start"]
        assert urgent["intent_type"] == "safety_rejected"
        assert "pre_safety" in urgent_stages and "inquiry" not in urgent_stages
        assert "episode_write" in urgent_stages

        async def empty_retrieval(plan):
            return [], {"warning": ""}
        monkeypatch.setattr(workflow, "retrieve_classics", empty_retrieval)
        exhausted = client.post("/api/chat", headers=headers,
                                json={"message": "我最近咳嗽三天，伴有疲劳"}).json()
        exhausted_trace = client.get(f'/api/threads/{exhausted["thread_id"]}/messages', headers=headers).json()[-1]["trace"]["events"]
        exhausted_stages = [event["stage"] for event in exhausted_trace if event["status"] == "start"]
        # All supported expansions already participated; do not repeat identical work.
        assert exhausted_stages.count("search") == 1
        assert exhausted_stages.count("rewrite_query") == 0
        assert "检索结果相关性不足" in exhausted["response"]
        assert "当前没有可引用的古籍片段" in exhausted["response"]
